fron PIL import Image
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt  # Pt(간격 조절용) 추가
import io
import re

# 파일 이름으로 사용할 수 없는 특수문자 제거 함수
def sanitize_filename(filename):
    return re.sub(r'[\\/*?:"<>|]', "", filename).strip()

st.set_page_config(page_title="현장 사진대지 생성기", layout="wide")
st.title("🏗️ 공사 항목별 사진대지 개별 생성")

if 'sections' not in st.session_state:
    st.session_state.sections = [{"title": "", "before": [], "after": []}]

def add_section():
    st.session_state.sections.append({"title": "", "before": [], "after": []})

def remove_section(index):
    if len(st.session_state.sections) > 1:
        st.session_state.sections.pop(index)

st.button("➕ 새로운 공정(항목) 추가", on_click=add_section)

for i, section in enumerate(st.session_state.sections):
    with st.container():
        st.markdown(f"### [항목 {i+1}]")
        col_title, col_del = st.columns([8, 1])
        with col_title:
            section['title'] = st.text_input(f"작업 명칭 입력", key=f"title_{i}", value=section['title'], placeholder="예: 정화조 철거, 내장재 해체 등")
        with col_del:
            if st.button("🗑️ 삭제", key=f"del_{i}"):
                remove_section(i)
                st.rerun()

        col_before, col_after = st.columns(2)
        with col_before:
            st.subheader("📷 '전' 사진")
            section['before'] = st.file_uploader(f"항목 {i+1} 전 사진", type=['jpg', 'png', 'jpeg'], accept_multiple_files=True, key=f"before_{i}")
        with col_after:
            st.subheader("📸 '후' 사진")
            section['after'] = st.file_uploader(f"항목 {i+1} 후 사진", type=['jpg', 'png', 'jpeg'], accept_multiple_files=True, key=f"after_{i}")
        st.divider()

def create_docx_content(mode="before"):
    doc = Document()
    suffix = " (작업 전)" if mode == "before" else " (철거 후)"
    doc.add_heading("공사 사진 대지" + suffix, 0)

    for section in st.session_state.sections:
        photos = section['before'] if mode == "before" else section['after']
        if not photos: continue
            
        current_title = section['title'] if section['title'] else "제목 없음"
        doc.add_heading(current_title, level=1)
        
        # 표 생성
        table = doc.add_table(rows=0, cols=1)
        table.style = 'Table Grid'
        
        # 헤더 (표 상단 제목)
        hdr = table.add_row().cells
        hdr[0].text = "촬영 사진 (" + ("전" if mode == "before" else "후") + ")"
        
        for photo in photos:
            row = table.add_row().cells
            paragraph = row[0].paragraphs[0]
            
            # --- 이미지 압축 로직 추가 ---
            # 1. 사용자가 올린 사진을 엽니다
            img = Image.open(photo)
            
            # 2. 사진 방향(회전) 정보가 있으면 올바르게 잡아줍니다 (스마트폰 사진 뒤집힘 방지)
            from PIL import ImageOps
            img = ImageOps.exif_transpose(img)
            
            # 3. 사진의 해상도를 가로 1000픽셀 수준으로 줄입니다 (웹/문서용으로 충분함)
            max_size = (1000, 1000)
            img.thumbnail(max_size, Image.Resampling.LANCZOS)
            
            # 4. 압축한 이미지를 임시 공간에 저장합니다
            img_io = io.BytesIO()
            # JPEG 형식으로 변환하여 저장 (투명 배경 있는 PNG 오류 방지)
            if img.mode in ("RGBA", "P"): 
                img = img.convert("RGB")
            img.save(img_io, format='JPEG', quality=85) # quality 85면 화질은 좋고 용량은 확 줄어듦
            img_io.seek(0)
            
            # 5. 압축된 이미지를 워드 문서에 넣습니다 (너비는 여전히 5.5인치 유지)
            run = paragraph.add_run()
            run.add_picture(img_io, width=Inches(5.5))
            
            # 여백 및 정렬
            paragraph.paragraph_format.space_after = Pt(35)
            paragraph.alignment = 1
            
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# --- 파일 이름 생성 로직 ---
titles = [s['title'] for s in st.session_state.sections if s['title'].strip()]
if not titles:
    base_name = "현장사진대지"
elif len(titles) == 1:
    base_name = sanitize_filename(titles[0])
else:
    base_name = f"{sanitize_filename(titles[0])}_외_{len(titles)-1}건"

file_name_before = f"{base_name}_작업전.docx"
file_name_after = f"{base_name}_작업후.docx"

st.write(f"### 🖨️ 문서 생성 (파일명: {base_name}...)")
btn_col1, btn_col2 = st.columns(2)

with btn_col1:
    before_data = create_docx_content("before")
    st.download_button(
        label="📄 [작업 전] 문서 생성 및 다운로드",
        data=before_data,
        file_name=file_name_before,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True
    )

with btn_col2:
    after_data = create_docx_content("after")
    st.download_button(
        label="📄 [작업 후] 문서 생성 및 다운로드",
        data=after_data,
        file_name=file_name_after,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True
    )